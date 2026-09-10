"""Build the review DOCX from ANALYST_GUIDE.md with python-docx.

Run with a document-tooling Python (requires python-docx), independent of the
application environment. The guide intentionally contains no tables.
Preset: compact_reference_guide. Header pattern: editorial_cover.
Named overrides: compact cover title, 9pt running furniture, 8pt reference URLs.
"""
from pathlib import Path
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUT = ROOT / "Maritime_Carbon_Analyst_Guide.docx"


def style(doc, name, size, before=0, after=6, color="202B33", bold=False):
    item = doc.styles[name]
    item.font.name = "Calibri"
    item.font.size = Pt(size)
    item.font.bold = bold
    item.font.color.rgb = RGBColor.from_string(color)
    item.paragraph_format.space_before = Pt(before)
    item.paragraph_format.space_after = Pt(after)
    item.paragraph_format.line_spacing = 1.25
    item.paragraph_format.widow_control = True
    item.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Calibri")
    return item


def build():
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Inches(1)
    section.header_distance = section.footer_distance = Inches(.492)
    style(doc, "Normal", 11)
    style(doc, "Title", 28, before=22, after=14, color="163638", bold=True)
    style(doc, "Subtitle", 15, after=16, color="53666A")
    for name, size, before, after, color in (("Heading 1", 16, 18, 10, "2E74B5"),
        ("Heading 2", 13, 14, 7, "2E74B5"), ("Heading 3", 12, 10, 5, "1F4D78")):
        item = style(doc, name, size, before, after, color, True)
        item.paragraph_format.keep_with_next = True
        item.paragraph_format.keep_together = True
    bullet = style(doc, "List Bullet", 11, after=4)
    bullet.paragraph_format.left_indent = Inches(.375)
    bullet.paragraph_format.first_line_indent = Inches(-.188)
    for name in ("Header", "Footer"):
        style(doc, name, 9, after=0, color="647478")
    # Explicit numbering geometry; no literal bullet characters in paragraphs.
    numbering = doc.part.numbering_part.element
    for abstract in numbering.findall(qn("w:abstractNum")):
        for level in abstract.findall(qn("w:lvl")):
            props = level.find(qn("w:pPr"))
            if props is None:
                props = OxmlElement("w:pPr"); level.append(props)
            indent = props.find(qn("w:ind"))
            if indent is None:
                indent = OxmlElement("w:ind"); props.append(indent)
            indent.set(qn("w:left"), "540")
            indent.set(qn("w:hanging"), "271")
            tabs = props.find(qn("w:tabs"))
            if tabs is None:
                tabs = OxmlElement("w:tabs"); props.append(tabs)
            for tab in list(tabs):
                tabs.remove(tab)
            tab = OxmlElement("w:tab"); tab.set(qn("w:val"), "num"); tab.set(qn("w:pos"), "540"); tabs.append(tab)
    section.header.paragraphs[0].text = "MARITIME CARBON  |  ANALYST OPERATING GUIDE"
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("Review draft · Engine 6.6.0  |  ")
    field = OxmlElement("w:fldSimple"); field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    doc.core_properties.title = "Maritime Carbon — Analyst operating guide"
    doc.core_properties.subject = "Desktop application field reference and analysis workflow"
    doc.core_properties.author = "Maritime Carbon"
    doc.core_properties.keywords = "FuelEU; EU ETS; analyst; guide; review draft"
    doc.core_properties.comments = "Prepared for owner review. Fictional worked examples."
    source = (ROOT / "ANALYST_GUIDE.md").read_text(encoding="utf-8")
    for block in source.split("\n\n"):
        block = block.strip()
        if not block:
            continue
        if block.startswith("# "):
            p = doc.add_paragraph(block[2:], "Title")
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif block == "## Analyst operating guide":
            p = doc.add_paragraph("Analyst operating guide", "Subtitle")
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif block.startswith("## "):
            p = doc.add_paragraph(block[3:], "Heading 1")
            # Start the handbook after its compact opening and reading route.
            if block.startswith(("## 1.", "## 12.")):
                p.paragraph_format.page_break_before = True
        elif block.startswith("### "):
            doc.add_paragraph(block[4:], "Heading 2")
        elif block.startswith("- "):
            for line in block.splitlines():
                p = doc.add_paragraph(line.removeprefix("- "), "List Bullet")
                # Long source links receive a reusable reference override.
                if "https://" in line:
                    for run in p.runs:
                        run.font.size = Pt(8)
        else:
            p = doc.add_paragraph(block.replace("\n", " "))
            if block.startswith("Review draft ·"):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_after = Pt(22)
                for run in p.runs:
                    run.font.size = Pt(10)
    assert not doc.tables
    doc.save(OUT)
    # Structural content check, independent of pagination.
    reopened = Document(OUT)
    assert not reopened.tables
    expected = [re.sub(r"^#{1,3} |^- ", "", line) for line in source.splitlines() if line.strip()]
    actual = [p.text for p in reopened.paragraphs]
    assert actual == expected, "Document content differs from reviewed source"
    print(f"Created {OUT}: {len(actual)} paragraphs; {len(source.split())} words; zero tables")


if __name__ == "__main__":
    build()
