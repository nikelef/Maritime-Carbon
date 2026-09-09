"""Excel import workbook and its embedded, shared filling instructions."""
from io import BytesIO
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName
from .engine.schema import default_fuels
from .row_imports import VESSEL_COLUMNS, CONSUMPTION_COLUMNS, LEGS

EXAMPLE_VESSEL = ("ATLAS", "Demo Atlas", "", "cargo", 25000, "HVO", 100)
EXAMPLE_ROWS = [("ATLAS",1,"EU/non-EU sea","HSFO",1000,0),
                ("ATLAS",1,"EU/non-EU sea","VLSFO",100,0),
                ("ATLAS",1,"EU/non-EU sea","MGO",50,0),
                ("ATLAS",1,"Intra-EU sea","HSFO",200,0),
                ("ATLAS",1,"Intra-EU sea","VLSFO",50,0),
                ("ATLAS",1,"Intra-EU sea","MGO",50,0),
                ("ATLAS",1,"EU berth","MGO",30,0)]
FIELD_HELP = {
    "vessel_id": "Your short ship identifier, 1–50 characters. Use it once in Vessels and repeat it exactly in Consumption. It links rows within this upload; it does not update an existing app vessel.",
    "name": "The vessel name shown in the app, 1–100 characters. Use a distinct vessel_id for every ship even if names are similar.",
    "imo": "Optional seven-digit IMO number, including its valid check digit. Leave blank for a planning vessel. Store as text; do not enter an invented IMO number.",
    "ship_type": "Choose cargo (including bulk carriers), tanker, container or passenger. Use the code in the drop-down, not the longer description.",
    "gt": "Gross tonnage as a whole number from 1 to 1,000,000. This is GT, not deadweight or fuel tonnes. The app uses it when assessing regulatory scope.",
    "alternative": "The one alternative fuel to evaluate for this vessel, such as HVO. Use a fuel code available in the target case. This nomination does not enable compatibility or confirm fuel eligibility.",
    "max_alternative_pct": "Maximum alternative sea-service energy share, 0–100. Enter 30 for 30%, or 100 to permit full replacement. This is a cap, not an instruction to buy that quantity and not a mass blend ratio.",
    "voyage": "Whole number 1–100. Use the same number for the sea and berth rows belonging to one voyage. Each vessel may have its own Voyage 1. All rows are annual input quantities for the case.",
    "leg": "Choose Intra-EU sea, EU/non-EU sea, Outside-EU sea, EU berth or Non-EU berth. Sea and berth are separate categories. Enter full physical fuel tonnes; the app applies coverage itself.",
    "fuel": "Current fuel consumed on this row. Use HSFO, VLSFO, MGO or another case-library code. MGO is the app's MGO / LSMGO entry; LSMGO is accepted and converted to MGO. Include current pilot fuel once in its actual row.",
    "tonnes": "Fuel mass consumed during the year on this vessel/voyage/category, from 0 to 10,000,000 tonnes. Enter numbers only, e.g. 1000 or 12.5. Do not enter energy, currency, a total for all fuels or coverage-adjusted tonnes.",
    "eu_berth_ops_pct": "Only for EU berth: the share of this row's fuel use that could be replaced by shore power, 0–100. Blank or 0 means none. Leave blank or 0 for every other category. OPS remains disabled until reviewed and enabled in the app.",
}
INSTRUCTIONS = [
    ("START HERE", "This workbook matches the app's consumption table: one row per vessel, voyage, operating category and current fuel. Only Vessels and Consumption are imported. Instructions, Example and Lists are reference sheets."),
    ("1. Enter vessel profiles", "Fill Vessels from row 2, one row per ship. Assign a unique vessel_id and complete the identity, alternative fuel and maximum alternative share. The input sheets are blank so the example cannot be imported accidentally."),
    ("2. Enter consumption", "Fill Consumption from row 2. Repeat the vessel_id from Vessels, select the voyage, category and current fuel, then enter annual tonnes. Multiple fuels on the same leg require separate rows."),
    ("3. Avoid double counting", "A vessel/voyage/category/fuel combination may occur only once. Combine repeated quantities in that row. Do not add total rows to either input sheet. Do not also enter HVO replacement tonnes unless HVO is genuinely part of current baseline consumption."),
    ("4. Save and upload", "Save as Excel Workbook (.xlsx). In the decision case, choose Import vessels, select this file and click Validate & import vessels. Import appends new vessels and keeps existing ones. Uploading it twice adds the ships twice; it is not an update operation."),
    ("5. Review in the app", "Open each imported vessel and check all rows. Alternative compatibility and shore-power readiness start unconfirmed. Review engine class, existing pilot fuel, efficiency, investment costs, eligibility evidence and ship-specific FuelEU history before comparison."),
    ("Limits and file structure", "Use at most 20 vessel rows and 2,000 consumption rows, within the company's remaining vessel limit. Each vessel supports at most 100 distinct voyage/fuel groups. The .xlsx file limit is 2 MB. Preserve sheet names and row-1 headers; do not add data columns or password protection."),
    ("Values and percentages", "Enter checked values only in the input sheets. Paste special > Values if copying formulas. Blank rows are ignored. Use percentages as numbers from 0 to 100: type 30, not 0.30 or 30%. Use the drop-downs or exact supported labels. Do not halve EU/non-EU tonnes."),
    ("Vessels — vessel_id", FIELD_HELP["vessel_id"]),
    ("Vessels — name and IMO", FIELD_HELP["name"]+" "+FIELD_HELP["imo"]),
    ("Vessels — ship_type", FIELD_HELP["ship_type"]),
    ("Vessels — gt", FIELD_HELP["gt"]),
    ("Vessels — alternative", FIELD_HELP["alternative"]),
    ("Vessels — max_alternative_pct", FIELD_HELP["max_alternative_pct"]),
    ("Consumption — vessel_id and voyage", "Repeat the exact vessel_id from Vessels. "+FIELD_HELP["voyage"]),
    ("Consumption — leg", FIELD_HELP["leg"]),
    ("Consumption — fuel", FIELD_HELP["fuel"]),
    ("Consumption — tonnes", FIELD_HELP["tonnes"]),
    ("Consumption — eu_berth_ops_pct", FIELD_HELP["eu_berth_ops_pct"]),
    ("What remains in the app", "Set reporting years, EUA and pooling prices, fuel prices, regulatory scope and pooling capacity in the decision case. These are not imported by this fleet workbook. Confirm FuelEU eligibility and ETS CO2 zero-rating separately in the fuel library."),
    ("Worked example", "Example shows one bulk carrier and Voyage 1: EU/non-EU sea 1,000 t HSFO + 100 t VLSFO + 50 t MGO; intra-EU sea 200 t HSFO + 50 t VLSFO + 50 t MGO; EU berth 30 t MGO. Total physical fuel is 1,480 t. HVO is the nominated replacement fuel; shore power is zero."),
    ("To try the example", "Copy Example's vessel row into Vessels row 2, then its seven consumption rows into Consumption rows 2–8. Copy values only, excluding headings and totals. For the video case, separately set 2026 only, EUA 85 and pooling 190 in the app, then complete the technical and fuel reviews."),
    ("CSV compatibility", "The optional row-based CSV uses one combined row per consumption entry: repeat identical vessel details each time. Its size limit is 512 KB. The old 11-column CSV remains supported up to 64 KB, but only that older format uses sea totals and voyage percentages."),
    ("Reference", "Maritime Carbon application input contract. https://github.com/nikelef/Maritime-Carbon — use the case's fuel library as the authority for available fuel codes. This workbook imports planning inputs; it does not certify compliance."),
]

INK, TEAL, GREEN, LIGHT = "102E32", "087F72", "C8EEA0", "F0F7F5"


def build_workbook():
    wb = Workbook()
    wb.remove(wb.active)
    for name in ("Instructions","Vessels","Consumption","Example","Lists"):
        ws = wb.create_sheet(name)
        ws.sheet_view.showGridLines = False
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth,ws.page_setup.fitToHeight = 1,0
    def header(ws,row,values):
        for col,value in enumerate(values,1):
            c = ws.cell(row,col,value)
            c.fill,c.font = PatternFill("solid",fgColor=INK),Font(name="Calibri",size=11,bold=True,color="FFFFFF")
            c.alignment = Alignment(wrap_text=True,vertical="center")
        ws.row_dimensions[row].height = 34
    def note(ws,row,title,body,last=7):
        ws.merge_cells(start_row=row,start_column=1,end_row=row,end_column=last)
        c = ws.cell(row,1,title)
        c.font,c.fill = Font(name="Calibri",size=12,bold=True,color=INK),PatternFill("solid",fgColor=GREEN)
        c.alignment = Alignment(vertical="center")
        ws.row_dimensions[row].height = 25
        ws.merge_cells(start_row=row+1,start_column=1,end_row=row+1,end_column=last)
        c = ws.cell(row+1,1,body)
        c.font,c.alignment = Font(name="Calibri",size=11,color=INK),Alignment(wrap_text=True,vertical="center")
        ws.row_dimensions[row+1].height = 48 if len(body)<300 else 64
    ws = wb["Instructions"]
    for col in range(1,8):
        ws.column_dimensions[get_column_letter(col)].width = 18
    note(ws,1,"MARITIME CARBON | FLEET IMPORT", "Fill Vessels and Consumption. Keep the instructions with the workbook. Blank input sheets are ready for your data; the separate Example sheet is not imported.")
    for i,(title,body) in enumerate(INSTRUCTIONS):
        note(ws,4+i*3,title,body)
    ws.freeze_panes,ws.print_area = "A4",f"A1:G{2+len(INSTRUCTIONS)*3}"
    ws.sheet_view.zoomScale = 85
    for name,columns,widths,count in (("Vessels",VESSEL_COLUMNS,[18,27,17,18,15,20,26],20),
                                    ("Consumption",CONSUMPTION_COLUMNS,[18,14,25,19,20,28],100)):
        ws = wb[name]
        header(ws,1,columns)
        for col,(key,width) in enumerate(zip(columns,widths),1):
            ws.column_dimensions[get_column_letter(col)].width = width
            ws.cell(1,col).comment = Comment(FIELD_HELP[key],"Maritime Carbon")
            for row in range(2,count+2):
                c = ws.cell(row,col)
                c.font = Font(name="Calibri",size=11,color="1764A0")
                c.fill = PatternFill("solid",fgColor=LIGHT if row%2==0 else "FFFFFF")
                c.alignment = Alignment(vertical="center")
                c.number_format = "@" if key in ("vessel_id","name","imo","ship_type","alternative","leg","fuel") else "0" if key in ("gt","voyage") else "0.000"
                ws.row_dimensions[row].height = 23
        ws.freeze_panes = "B2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{21 if name=='Vessels' else 2001}"
        ws.print_title_rows = "1:1"
        ws.print_area = f"A1:{get_column_letter(len(columns))}{count+1}"
    ws = wb["Lists"]
    header(ws,1,["Ship type code","Voyage / berth category","Fuel code","Fuel description"])
    for col,width in enumerate([22,27,22,65],1):
        ws.column_dimensions[get_column_letter(col)].width = width
    for row,code in enumerate(("cargo","tanker","container","passenger"),2):
        ws.cell(row,1,code)
    for row,leg in enumerate(LEGS,2):
        ws.cell(row,2,leg)
    for row,f in enumerate(default_fuels(),2):
        ws.cell(row,3,f.code)
        ws.cell(row,4,"MGO / LSMGO · marine gas oil" if f.code=="MGO" else f.name)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = Font(name="Calibri",size=11,color=INK)
            c.alignment = Alignment(wrap_text=True,vertical="center")
        ws.row_dimensions[row[0].row].height = 26
    for name,ref in {"ShipTypes":"'Lists'!$A$2:$A$5","LegLabels":"'Lists'!$B$2:$B$6",
                     "FuelCodes":f"'Lists'!$C$2:$C${len(default_fuels())+1}","VesselIDs":"'Vessels'!$A$2:$A$21"}.items():
        wb.defined_names.add(DefinedName(name,attr_text=ref))
    def validation(sheet,area,kind,formula1,formula2=None):
        dv = DataValidation(type=kind,formula1=formula1,formula2=formula2,allow_blank=True)
        if kind!="list":
            dv.operator = "between"
        dv.showErrorMessage,dv.errorTitle,dv.error = True,"Check this input","Use a listed value or a number within the permitted range. See Instructions."
        wb[sheet].add_data_validation(dv)
        dv.add(area)
    for sheet,area,kind,first,last in [("Vessels","D2:D21","list","ShipTypes",None),("Vessels","F2:F21","list","FuelCodes",None),
        ("Vessels","E2:E21","whole",1,1_000_000),("Vessels","G2:G21","decimal",0,100),
        ("Consumption","A2:A2001","list","VesselIDs",None),("Consumption","B2:B2001","whole",1,100),
        ("Consumption","C2:C2001","list","LegLabels",None),("Consumption","D2:D2001","list","FuelCodes",None),
        ("Consumption","E2:E2001","decimal",0,10_000_000),("Consumption","F2:F2001","decimal",0,100)]:
        validation(sheet,area,kind,first,last)
    ws = wb["Example"]
    for col,width in enumerate([18,27,25,23,18,22,26],1):
        ws.column_dimensions[get_column_letter(col)].width = width
    note(ws,1,"ILLUSTRATIVE EXAMPLE | NOT IMPORTED", "Copy row 5 into Vessels row 2 and rows 9–15 into Consumption rows 2–8, using values only. Reporting years and prices are set separately in the app.")
    header(ws,4,VESSEL_COLUMNS)
    for col,value in enumerate(EXAMPLE_VESSEL,1):
        ws.cell(5,col,value)
    header(ws,8,CONSUMPTION_COLUMNS)
    for row,values in enumerate(EXAMPLE_ROWS,9):
        for col,value in enumerate(values,1):
            ws.cell(row,col,value)
    ws.cell(17,4,"Total fuel · tonnes")
    ws.cell(17,5,"=SUM(E9:E15)")
    note(ws,19,"READ WITH THE EXAMPLE", "HVO is nominated as the alternative; it is not added to current consumption. MGO represents LSMGO. All berth fuel is MGO and shore-power shares are zero. The example uses full physical tonnes, not tonnes reduced for 50% voyage coverage.")
    for row in list(ws.iter_rows(min_row=5,max_row=5))+list(ws.iter_rows(min_row=9,max_row=17)):
        for c in row:
            c.font = Font(name="Calibri",size=11,color=INK)
            c.alignment = Alignment(vertical="center")
            if isinstance(c.value,(float,int)) or c.data_type=="f":
                c.number_format = "#,##0.000" if c.column==5 and c.row>=9 else "#,##0"
        ws.row_dimensions[row[0].row].height = 25
    ws.print_area = "A1:G20"
    wb.active = 0
    return wb


def template_bytes():
    stream = BytesIO()
    build_workbook().save(stream)
    return stream.getvalue()
