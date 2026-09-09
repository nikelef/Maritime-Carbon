"""Read explicit voyage/leg/fuel inputs from the Excel template or matching CSV."""
import csv
import io
import math
from zipfile import ZipFile, BadZipFile
from xml.etree.ElementTree import ParseError
from defusedxml.common import DefusedXmlException
from openpyxl import load_workbook
from .engine.schema import Vessel

VESSEL_COLUMNS = ("vessel_id", "name", "imo", "ship_type", "gt", "alternative", "max_alternative_pct")
CONSUMPTION_COLUMNS = ("vessel_id", "voyage", "leg", "fuel", "tonnes", "eu_berth_ops_pct")
ROW_COLUMNS = VESSEL_COLUMNS + CONSUMPTION_COLUMNS[1:]
LEGS = {"Intra-EU sea": "intra_t", "EU/non-EU sea": "extra_t", "Outside-EU sea": "outside_t",
        "EU berth": "eu_berth_t", "Non-EU berth": "non_eu_berth_t"}
MAX_ROWS = 2000


def text(value):
    if value is None:
        return ""
    if isinstance(value,float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def number(value, field, where, low, high, integer=False):
    try:
        if isinstance(value,bool):
            raise ValueError()
        result = float(value)
        if not math.isfinite(result) or not low<=result<=high or (integer and not result.is_integer()):
            raise ValueError()
    except (TypeError,ValueError,OverflowError) as exc:
        raise ValueError(f"{where}: {field} must be a {'whole' if integer else 'finite'} number from {low:g} to {high:g}.") from exc
    return int(result) if integer else result


def fuel(value):
    code = text(value)
    return "MGO" if code.upper() in ("LSMGO","MGO / LSMGO") else code


def parse_rows(vessel_rows, consumption_rows):
    vessels,groups,seen = {},{},set()
    for where,row in vessel_rows:
        identity = text(row["vessel_id"])
        if not identity or len(identity)>50:
            raise ValueError(f"{where}: enter a vessel_id of 1–50 characters.")
        if identity in vessels:
            raise ValueError(f"{where}: duplicate vessel_id {identity}; use one Vessels row per ship.")
        if len(vessels)>=20:
            raise ValueError("Import at most 20 vessels at a time.")
        vessels[identity] = Vessel(name=text(row["name"]),imo=text(row["imo"]),ship_type=text(row["ship_type"]),
            gt=number(row["gt"],"gt",where,1,1_000_000,True),alternative=fuel(row["alternative"]),
            max_blend=number(row["max_alternative_pct"],"max_alternative_pct",where,0,100)/100,
            sea_t=0,berth_t=0,intra_share=0,extra_share=0,ops_fraction=0,consumption=[])
        groups[identity] = {}
    if not vessels:
        raise ValueError("The Vessels sheet has no vessel rows. Enter your ships before uploading.")
    for where,row in consumption_rows:
        identity = text(row["vessel_id"])
        if identity not in vessels:
            raise ValueError(f"{where}: vessel_id {identity or '(blank)'} is missing from Vessels.")
        voyage = number(row["voyage"],"voyage",where,1,100,True)
        leg,code = text(row["leg"]),fuel(row["fuel"])
        if leg not in LEGS:
            raise ValueError(f"{where}: choose a supported leg: {', '.join(LEGS)}.")
        if not code:
            raise ValueError(f"{where}: enter a fuel code from this case's fuel library.")
        key = (identity,voyage,leg,code)
        if key in seen:
            raise ValueError(f"{where}: duplicate vessel/voyage/leg/fuel row. Combine its tonnes into one row.")
        seen.add(key)
        tonnes = number(row["tonnes"],"tonnes",where,0,10_000_000)
        ops = number(row["eu_berth_ops_pct"] if text(row["eu_berth_ops_pct"]) else 0,"eu_berth_ops_pct",where,0,100)/100
        if leg!="EU berth" and ops:
            raise ValueError(f"{where}: shore-power share applies only to EU berth rows; leave it blank or zero.")
        bucket = groups[identity]
        group = bucket.setdefault((voyage,code),dict(voyage=voyage,fuel=code))
        if len(bucket)>100:
            raise ValueError(f"{where}: import at most 100 voyage/fuel groups per vessel.")
        group[LEGS[leg]] = tonnes
        if leg=="EU berth":
            group["ops_share"] = ops
    for identity,vessel in vessels.items():
        vessel.consumption = list(groups[identity].values())
        if not vessel.consumption:
            raise ValueError(f"Vessel {identity}: enter at least one Consumption row.")
    return list(vessels.values())


def sheet_rows(sheet, columns, limit):
    if (sheet.max_column or 0)>len(columns) or (sheet.max_row or 0)>limit+1:
        raise ValueError(f"{sheet.title}: preserve the template columns and use at most {limit} data rows.")
    rows = sheet.iter_rows(max_row=limit+2,max_col=len(columns))
    if tuple(text(c.value) for c in next(rows))!=columns:
        raise ValueError(f"{sheet.title}: headers must match the Excel template exactly in row 1.")
    result = []
    for index,cells in enumerate(rows,start=2):
        if all(c.value in (None,"") for c in cells):
            continue
        if index>limit+1:
            raise ValueError(f"{sheet.title}: use at most {limit} data rows.")
        if any(c.data_type in ("f","e") for c in cells):
            raise ValueError(f"{sheet.title} row {index}: use values only; replace formulas and Excel errors with checked values.")
        result.append((f"{sheet.title} row {index}",dict(zip(columns,(c.value for c in cells)))))
    return result


def parse_excel(content):
    workbook = None
    try:
        with ZipFile(io.BytesIO(content)) as archive:
            members = archive.infolist()
            if len(members)>250 or sum(x.file_size for x in members)>16*1024*1024:
                raise ValueError("The workbook is too large when expanded. Use a fresh import template.")
            if any(x.flag_bits & 1 for x in members):
                raise ValueError("Remove workbook password protection before importing.")
            if any("vbaproject" in x.filename.lower() or "externallinks/" in x.filename.lower() for x in members):
                raise ValueError("Use a plain .xlsx workbook without macros or external workbook links.")
        workbook = load_workbook(io.BytesIO(content),read_only=True,data_only=False,keep_links=False)
        if not {"Vessels","Consumption"}.issubset(workbook.sheetnames):
            raise ValueError("The workbook must contain the Vessels and Consumption sheets from the template.")
        return parse_rows(sheet_rows(workbook["Vessels"],VESSEL_COLUMNS,20),
                          sheet_rows(workbook["Consumption"],CONSUMPTION_COLUMNS,MAX_ROWS))
    except (BadZipFile,KeyError,ParseError,DefusedXmlException,OSError,TypeError,OverflowError) as exc:
        raise ValueError("The Excel file could not be read. Save a fresh .xlsx template and try again.") from exc
    finally:
        if workbook is not None:
            workbook.close()


def parse_row_csv(content):
    try:
        reader = csv.DictReader(io.StringIO(content.decode("utf-8-sig")),strict=True)
        if reader.fieldnames is None or len(reader.fieldnames)!=len(ROW_COLUMNS) or set(reader.fieldnames)!=set(ROW_COLUMNS):
            raise ValueError("CSV headers must match a downloadable template exactly.")
        vessels,consumption = {},[]
        for index,row in enumerate(reader,start=2):
            where = f"Row {index}"
            if index>MAX_ROWS+1:
                raise ValueError("Import at most 2000 consumption rows.")
            if None in row or any(v is None for v in row.values()):
                raise ValueError(f"{where}: wrong number of columns.")
            identity = row["vessel_id"].strip()
            metadata = {c:row[c].strip() for c in VESSEL_COLUMNS}
            if identity in vessels and vessels[identity][1]!=metadata:
                raise ValueError(f"{where}: repeat identical vessel details for vessel_id {identity}.")
            vessels.setdefault(identity,(where,metadata))
            consumption.append((where,{c:row[c] for c in CONSUMPTION_COLUMNS}))
        return parse_rows(vessels.values(),consumption)
    except (UnicodeDecodeError,csv.Error) as exc:
        raise ValueError("Save the file as CSV UTF-8 using the template and try again.") from exc
