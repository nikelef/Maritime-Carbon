"""Workbook download/upload, row accounting and all-or-nothing validation."""
import csv
import io
from datetime import timedelta
from unittest.mock import patch
from zipfile import ZipFile
from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook
from portal.engine.schema import ScenarioInput
from portal.imports import parse_vessels
from portal.import_template import build_workbook, EXAMPLE_VESSEL, EXAMPLE_ROWS
from portal.row_imports import ROW_COLUMNS
from portal.models import Company, Membership, Scenario


def example_book():
    book = build_workbook()
    for col,value in enumerate(EXAMPLE_VESSEL,1):
        book["Vessels"].cell(2,col,value)
    for row,values in enumerate(EXAMPLE_ROWS,2):
        for col,value in enumerate(values,1):
            book["Consumption"].cell(row,col,value)
    return book


def upload(book):
    stream = io.BytesIO()
    book.save(stream)
    return SimpleUploadedFile("fleet.xlsx",stream.getvalue(),"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@override_settings(PASSWORD_HASHERS=["django.contrib.auth.hashers.MD5PasswordHasher"])
class ExcelImportTests(TestCase):
    def setUp(self):
        self.company = Company.objects.create(name="Excel importer",status="active",access_until=timezone.localdate()+timedelta(days=30))
        self.user = get_user_model().objects.create_user("excel_user",password="test-password")
        Membership.objects.create(company=self.company,user=self.user,role="analyst")
        self.case = Scenario.objects.create(company=self.company,created_by=self.user,name="Row case",inputs=ScenarioInput().to_dict())
        self.client.force_login(self.user)

    def post(self,book,revision=1):
        return self.client.post(reverse("vessel_import",args=[self.case.pk]),{"revision":revision,"csv_file":upload(book)})

    def assert_unchanged(self):
        self.case.refresh_from_db()
        self.assertEqual(self.case.inputs["vessels"],[])
        self.assertEqual(self.case.revision,1)

    def test_download_has_instructions_blank_inputs_example_and_validation(self):
        response = self.client.get(reverse("vessel_template"))
        self.assertEqual(response.status_code,200)
        self.assertIn(".xlsx",response["Content-Disposition"])
        book = load_workbook(io.BytesIO(response.content))
        self.assertEqual(book.sheetnames,["Instructions","Vessels","Consumption","Example","Lists"])
        self.assertIsNone(book["Vessels"]["A2"].value)
        self.assertIsNone(book["Consumption"]["A2"].value)
        self.assertEqual(book["Example"]["E17"].value,"=SUM(E9:E15)")
        self.assertTrue(book["Consumption"].data_validations.dataValidation)
        self.assertIn("100",book["Vessels"]["G1"].comment.text)
        with self.assertRaisesRegex(ValueError,"no vessel rows"):
            parse_vessels(SimpleUploadedFile("blank.xlsx",response.content))

    def test_exact_video_example_round_trips_to_editor(self):
        book = example_book()
        book["Consumption"]["D4"] = "LSMGO"
        response = self.post(book)
        self.assertEqual(response.status_code,302)
        self.case.refresh_from_db()
        vessel = ScenarioInput.from_dict(self.case.inputs).vessels[0]
        self.assertEqual(vessel.total_sea_t,1450)
        self.assertEqual(vessel.total_eu_berth_t,30)
        self.assertEqual(vessel.total_non_eu_berth_t,0)
        self.assertEqual(len(vessel.consumption),3)
        mgo = next(r for r in vessel.consumption if r["fuel"]=="MGO")
        self.assertEqual(mgo,{"voyage":1,"fuel":"MGO","extra_t":50,"intra_t":50,"eu_berth_t":30,"ops_share":0})
        self.assertFalse(vessel.compatible)
        self.assertFalse(vessel.ops_ready)
        self.assertEqual(vessel.max_blend,1)
        editor = self.client.get(reverse("vessel_edit",args=[self.case.pk,0]))
        self.assertContains(editor,'value="1000.0"')
        self.assertContains(editor,'id_consumption-6-voyage')

    def test_multivessel_multivoyage_all_categories_preserve_quantities(self):
        book = example_book()
        for col,value in enumerate(("BETA","Second ship","","tanker",35000,"HVO",25),1):
            book["Vessels"].cell(3,col,value)
        additions = [("BETA",2,"Outside-EU sea","VLSFO",400,0),("BETA",2,"Non-EU berth","MGO",20,0),
                     ("BETA",3,"EU berth","MGO",10,50),("ATLAS",2,"Intra-EU sea","HSFO",12,0)]
        for row,values in enumerate(additions,9):
            for col,value in enumerate(values,1):
                book["Consumption"].cell(row,col,value)
        self.assertEqual(self.post(book).status_code,302)
        self.case.refresh_from_db()
        ships = ScenarioInput.from_dict(self.case.inputs).vessels
        self.assertEqual(len(ships),2)
        self.assertEqual(ships[0].total_sea_t,1462)
        self.assertEqual(ships[1].total_non_eu_berth_t,20)
        self.assertEqual(ships[1].consumption[2]["ops_share"],.5)
        self.assertFalse(ships[1].ops_ready)

    def test_bad_rows_are_atomic_and_explained(self):
        cases = [("A3","UNKNOWN","missing from Vessels"),("B3",1.5,"whole number"),
                 ("C3","Moon","supported leg"),("E3","NaN","finite number"),
                 ("E3","=10*10","values only"),("E3",-1,"finite number"),
                 ("F3",25,"only to EU berth"),("D3","NOT_A_FUEL","fuel library")]
        for cell,value,message in cases:
            with self.subTest(cell=cell,value=value):
                book = example_book()
                book["Consumption"][cell] = value
                self.assertContains(self.post(book),message)
                self.assert_unchanged()

    def test_duplicate_consumption_and_duplicate_vessel_rejected(self):
        for sheet,source,target in [("Vessels",2,3),("Consumption",2,9)]:
            book = example_book()
            for c in book[sheet][source]:
                book[sheet].cell(target,c.column,c.value)
            self.assertContains(self.post(book),"duplicate")
            self.assert_unchanged()

    def test_company_capacity_and_stale_revision_leave_existing_case_intact(self):
        self.assertContains(self.post(example_book(),99),"changed in another session")
        self.assert_unchanged()
        self.assertEqual(self.post(example_book()).status_code,302)
        self.company.max_vessels = 1
        self.company.save()
        self.assertContains(self.post(example_book(),2),"vessel limit")
        self.case.refresh_from_db()
        self.assertEqual(len(self.case.inputs["vessels"]),1)
        self.assertEqual(self.case.revision,2)

    def test_row_csv_download_matches_excel_and_requires_consistent_identity(self):
        response = self.client.get(reverse("vessel_template_csv"))
        vessels = parse_vessels(SimpleUploadedFile("rows.csv",response.content))
        self.assertEqual(vessels[0].consumption,parse_vessels(upload(example_book()))[0].consumption)
        rows = list(csv.reader(io.StringIO(response.content.decode())))
        self.assertEqual(tuple(rows[0]),ROW_COLUMNS)
        rows[2][1] = "Different name"
        stream = io.StringIO()
        csv.writer(stream).writerows(rows)
        with self.assertRaisesRegex(ValueError,"identical vessel details"):
            parse_vessels(SimpleUploadedFile("rows.csv",stream.getvalue().encode()))

    def test_file_limits_headers_and_formula_errors(self):
        for content,name in [(b"broken","bad.xlsx"),(b"x"*(2*1024*1024+1),"big.xlsx"),(b"x","old.xls")]:
            with self.assertRaises(ValueError):
                parse_vessels(SimpleUploadedFile(name,content))
        book = example_book()
        book["Vessels"]["A1"] = "company_id"
        with self.assertRaisesRegex(ValueError,"headers"):
            parse_vessels(upload(book))
        book = example_book()
        book["Consumption"]["E3"] = "#DIV/0!"
        with self.assertRaisesRegex(ValueError,"values only"):
            parse_vessels(upload(book))
        book = example_book()
        book["Consumption"]["A2002"] = "ATLAS"
        with self.assertRaisesRegex(ValueError,"2000"):
            parse_vessels(upload(book))

    def test_oversized_expansion_rejected_before_excel_parser(self):
        stream = io.BytesIO()
        with ZipFile(stream,"w") as archive:
            archive.writestr("xl/large.xml",b"x"*(16*1024*1024+1),compress_type=8)
        with patch("portal.row_imports.load_workbook") as loader,self.assertRaisesRegex(ValueError,"expanded"):
            parse_vessels(SimpleUploadedFile("large.xlsx",stream.getvalue()))
        loader.assert_not_called()
