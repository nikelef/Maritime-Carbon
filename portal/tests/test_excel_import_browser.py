"""Real Excel download → filled workbook upload → seven editor rows."""
import io
import os
from datetime import timedelta
from unittest import skipUnless
from django.conf import settings
from django.contrib.auth import get_user_model
from django.contrib.staticfiles.testing import StaticLiveServerTestCase
from django.test import override_settings
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook
from portal.engine.schema import ScenarioInput
from portal.import_template import EXAMPLE_VESSEL,EXAMPLE_ROWS
from portal.models import Company,Membership,Scenario


@skipUnless(os.environ.get("RUN_BROWSER_TESTS")=="1","Enable RUN_BROWSER_TESTS=1 for browser checks.")
@override_settings(STORAGES={"default":{"BACKEND":"django.core.files.storage.FileSystemStorage"},"staticfiles":{"BACKEND":"django.contrib.staticfiles.storage.StaticFilesStorage"}})
class ExcelBrowserTests(StaticLiveServerTestCase):
    def test_download_fill_upload(self):
        from playwright.sync_api import sync_playwright,expect
        company = Company.objects.create(name="Excel demo",status="active",access_until=timezone.localdate()+timedelta(days=30))
        user = get_user_model().objects.create_user("excel_browser")
        Membership.objects.create(company=company,user=user,role="analyst")
        case = Scenario.objects.create(company=company,created_by=user,name="Import demo",inputs=ScenarioInput().to_dict())
        self.client.force_login(user)
        with sync_playwright() as driver:
            browser = driver.chromium.launch(headless=True,channel=os.environ.get("PLAYWRIGHT_CHANNEL") or None)
            context = browser.new_context(viewport={"width":1600,"height":900})
            context.add_cookies([dict(name=settings.SESSION_COOKIE_NAME,value=self.client.cookies[settings.SESSION_COOKIE_NAME].value,url=self.live_server_url)])
            page = context.new_page()
            errors = []
            page.on("pageerror",lambda exc:errors.append(str(exc)))
            page.goto(self.live_server_url+reverse("vessel_import",args=[case.pk]))
            with page.expect_download() as download:
                page.get_by_role("link",name="Download Excel template").click()
            with open(download.value.path(),"rb") as downloaded:
                book = load_workbook(io.BytesIO(downloaded.read()))
            self.assertEqual(book.sheetnames[0],"Instructions")
            for col,value in enumerate(EXAMPLE_VESSEL,1):
                book["Vessels"].cell(2,col,value)
            for row,values in enumerate(EXAMPLE_ROWS,2):
                for col,value in enumerate(values,1):
                    book["Consumption"].cell(row,col,value)
            stream = io.BytesIO()
            book.save(stream)
            page.locator('input[type="file"]').set_input_files(dict(name="completed.xlsx",mimeType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",buffer=stream.getvalue()))
            page.get_by_role("button",name="Validate & import vessels").click()
            expect(page.get_by_role("heading",name="Every ship keeps its own balance.")).to_be_visible()
            page.get_by_role("link",name="Edit",exact=True).click()
            expect(page.locator(".consumption-row")).to_have_count(7)
            expect(page.locator("#id_alternative")).to_have_value("HVO")
            self.assertEqual(errors,[])
            browser.close()
        case.refresh_from_db()
        self.assertEqual(ScenarioInput.from_dict(case.inputs).vessels[0].total_sea_t,1450)
